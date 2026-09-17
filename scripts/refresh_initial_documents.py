"""Refresh the archived A-C coursework, keeping editorial records outside submissions."""
from pathlib import Path
from docx import Document
from copy import deepcopy
import json

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / 'archive/2026-09-17-before-reorganization/ClassMic_C1_Completed'
OUT = ROOT / 'submissions/initial-phases'
ITEMS = [
    ('Supporting_A2_B2/ClassMic_A2_Aspiration.docx','A_Aspiration/ClassMic_A2_Aspiration.docx'),
    ('Supporting_A2_B2/ClassMic_B2_Business_Case.docx','B_Business_Case/ClassMic_B2_Business_Case.docx'),
    ('ClassMic_C1_Project_Charter.docx','C_Charter/ClassMic_C1_Project_Charter.docx'),
]
REPLACEMENTS = {
 '10 September 2026':'17 September 2026',
 'by 19 September':'by 18 September',
 'By 19 September':'By 18 September',
 'for 19 September':'for 18 September',
 '10–19 September':'10–18 September',
 '20–26 September':'19–25 September',
 '27 September–3 October':'26 September–2 October',
 '4–9 October':'3–9 October',
 'Draft project charter. ClassMic proposes a two-sprint browser prototype and controlled evaluation, based on the project aspiration and business case.':'Proposed project charter. ClassMic will test a browser microphone workflow through two weekly development sprints and a controlled evaluation.',
 'B2 recommends a small course prototype using existing resources.':'The proposed course project uses existing resources for a small prototype.',
 'ClassMic aspiration and charter':'Proposed course prototype scope',
 'Formal course acceptance remains open.':'The instructor reviews the course proposal.',
 'Course acceptance authority: the BUS 2010 instructor. Formal charter approval and the named sponsor record remain open.':'Course acceptance authority: the BUS 2010 instructor. The proposed charter takes effect after course approval.',
 'These are the selected planning roles.':'These responsibilities form the proposed staffing plan.',
 'accepts internal work against the PRD':'reviews work against the agreed requirements',
 'The project concept comes from Temiko’s in-class SMART note. Course milestones follow the published BUS 2010 LMS schedule.':'The concept originated in classroom planning. The proposed schedule follows the BUS 2010 course deadlines.',
 'One supplied in-class SMART note':'The initial classroom concept',
 'Temiko’s in-class SMART note':'The initial classroom concept',
 'The in-class SMART note proposes':'The proposed concept uses',
 'In-class SMART note':'Initial classroom concept',
 'original SMART note':'initial classroom concept',
}

def paragraphs(d):
    yield from d.paragraphs
    for t in d.tables:
        for row in t.rows:
            for c in row.cells:
                yield from c.paragraphs
    for s in d.sections:
        yield from s.header.paragraphs
        yield from s.footer.paragraphs

for source,target in ITEMS:
    d=Document(SOURCE/source)
    # Remove preparation-only checklist; real assumptions remain in relevant sections.
    deleting=False
    for p in list(d.paragraphs):
        if p.text.strip()=='Items requiring confirmation': deleting=True
        elif deleting and (p.text.strip()=='Sources' or p.text.startswith('[1]')): deleting=False
        if deleting: p._element.getparent().remove(p._element)
    for p in paragraphs(d):
        text=p.text
        for old,new in REPLACEMENTS.items(): text=text.replace(old,new)
        if text != p.text:
            if p.runs:
                p.runs[0].text=text
                for run in p.runs[1:]: run.text=''
            else: p.add_run(text)
    if 'Business_Case' in target:
        for p in d.paragraphs:
            if p.text.startswith('Planning review is scheduled'):
                p.text='Planning review is scheduled for 18 September, before two full weekly build sprints. The planning cutoff is 22 September at 00:00. The final product report is due 3 October at 00:00, so prototype development and evaluation must finish on 2 October. Execution evidence is due 4 October at 00:00 and closing work on 10 October at 00:00. These are proposed internal work dates, subject to available team capacity.'
    if 'Charter' in target:
        for p in d.paragraphs:
            if p.text.startswith('Course milestones follow'):
                p.text='Course cutoffs are 22 September for planning, 3 October for the final product report, 4 October for execution evidence and 10 October for closing, all at 00:00. Complete the product report by the evening of 2 October.'
    if 'Aspiration' in target:
        for p in d.paragraphs:
            if p.text.startswith('[1]'):
                p.text='[1] Biamp Crowd Mics product overview, accessed 17 September 2026. https://www.biamp.com/products/families/crowd-mics'
    edits=json.loads((ROOT/'scripts/editorial_copy.json').read_text())['documents'][Path(target).stem]
    seen=set()
    for para in paragraphs(d):
        old=para.text
        if old not in edits: continue
        seen.add(old)
        runs=[run for run in para.runs if run.text]
        properties=deepcopy(runs[0]._r.rPr) if runs and runs[0]._r.rPr is not None else None
        mixed_bold=any(run.bold is True for run in runs) and not all(run.bold is True for run in runs)
        para.clear()
        run=para.add_run(edits[old])
        if properties is not None: run._r.insert(0,properties)
        if mixed_bold: run.bold=False
    assert seen==set(edits), 'Unmatched editorial replacements: '+str(set(edits)-seen)
    for section in d.sections:
        for para in section.footer.paragraphs:
            for run in para.runs: run.text=run.text.replace('C1 draft for review','C1')
    d.core_properties.author='Temiko Machavariani and Nurtore Arynuruly'
    d.core_properties.last_modified_by=''
    d.core_properties.comments=''
    d.core_properties.keywords=''
    d.core_properties.category=''
    d.save(OUT/target)
    print(target)
