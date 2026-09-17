from pathlib import Path
from datetime import date
import csv, json, shutil
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT

ROOT = Path('/Users/gravitycodez/Desktop/Project managment')
OUT = ROOT / 'ClassMic_Submission'
for folder in ['A_Aspiration', 'B_Business_Case', 'C_Charter', 'D_Develop_Plans', 'Supporting_Notes']:
    (OUT / folder).mkdir(parents=True, exist_ok=True)

TEAM = 'Temiko Machavariani and Nurtore Arynuruly'
BIO = 'https://www.biamp.com/products/families/crowd-mics'
FAQ = 'https://www.biamp.com/products/families/crowd-mics/faq'
AV = 'https://support.biamp.com/CrowdMics/Crowd_Mics_design_guide'
W3 = 'https://www.w3.org/TR/mediacapture-streams/'

def base(title, label, body_size=11):
    d = Document()
    for border in list(d.styles.element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
    s = d.sections[0]
    s.page_width, s.page_height = Inches(8.5), Inches(11)
    s.top_margin = s.bottom_margin = Inches(.65)
    s.left_margin = s.right_margin = Inches(.7)
    s.header_distance = s.footer_distance = Inches(.28)
    for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2']:
        st = d.styles[name]
        st.font.name = 'Arial'
        st.font.color.rgb = RGBColor(0, 0, 0)
        st.paragraph_format.line_spacing = 1.04
    normal = d.styles['Normal']
    normal.font.size = Pt(body_size)
    normal.paragraph_format.space_after = Pt(6)
    title_st = d.styles['Title']
    title_st.font.size = Pt(23)
    title_st.font.bold = True
    title_st.paragraph_format.space_after = Pt(5)
    d.styles['Subtitle'].font.size = Pt(11)
    d.styles['Subtitle'].font.italic = False
    d.styles['Subtitle'].paragraph_format.space_after = Pt(8)
    for name in ['Heading 1', 'Heading 2']:
        st = d.styles[name]
        st.font.size = Pt(12)
        st.font.bold = True
        st.paragraph_format.space_before = Pt(9)
        st.paragraph_format.space_after = Pt(5)
        st.paragraph_format.keep_with_next = True
    f = s.footer.paragraphs[0]
    f.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = f.add_run('ClassMic   ' + label + '   ')
    r.font.size = Pt(8)
    fld = OxmlElement('w:fldSimple'); fld.set(qn('w:instr'), 'PAGE'); f._p.append(fld)
    d.core_properties.title = title
    d.core_properties.subject = 'BUS 2010 ClassMic course project'
    d.core_properties.author = TEAM
    d.core_properties.last_modified_by = ''
    return d

def p(d, text, lead=None, size=None):
    x = d.add_paragraph()
    if lead and text.startswith(lead):
        x.add_run(lead).bold = True
        x.add_run(text[len(lead):])
    else:
        x.add_run(text)
    if size:
        for r in x.runs: r.font.size = Pt(size)
    return x

def h(d, text): return d.add_paragraph(text, 'Heading 1')
def page(d, title):
    d.add_page_break()
    h(d, title)

def title(d, text, subtitle):
    d.add_paragraph(text, 'Title')
    d.add_paragraph(subtitle, 'Subtitle')
    p(d, TEAM + '   |   BUS 2010   |   10 September 2026', size=10)

def table(d, headers, rows, widths, font=10.5, center_cols=()):
    t = d.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for i, w in enumerate(widths): t.columns[i].width = Inches(w)
    for i, text in enumerate(headers): t.rows[0].cells[i].text = str(text)
    for row in rows:
        for c, text in zip(t.add_row().cells, row): c.text = str(text)
    borders = OxmlElement('w:tblBorders')
    for edge in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        e = OxmlElement('w:' + edge)
        for k, v in [('val', 'single'), ('sz', '4'), ('color', 'D9D9D9')]: e.set(qn('w:' + k), v)
        borders.append(e)
    t._tbl.tblPr.append(borders)
    for ri, row in enumerate(t.rows):
        pr = row._tr.get_or_add_trPr()
        pr.append(OxmlElement('w:cantSplit'))
        if ri == 0: pr.append(OxmlElement('w:tblHeader'))
        for ci, cell in enumerate(row.cells):
            cell.width = Inches(widths[ci])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            props = cell._tc.get_or_add_tcPr()
            shade = OxmlElement('w:shd')
            shade.set(qn('w:fill'), '263F50' if ri == 0 else ('F1F4F6' if ri % 2 else 'FFFFFF'))
            props.append(shade)
            margins = OxmlElement('w:tcMar')
            for side, val in [('top', 70), ('bottom', 70), ('left', 90), ('right', 90)]:
                e = OxmlElement('w:' + side); e.set(qn('w:w'), str(val)); e.set(qn('w:type'), 'dxa'); margins.append(e)
            props.append(margins)
            for par in cell.paragraphs:
                par.paragraph_format.space_after = Pt(0)
                par.paragraph_format.line_spacing = 1.02
                if ci in center_cols: par.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for r in par.runs:
                    r.font.name = 'Arial'; r.font.size = Pt(font)
                    r.font.bold = ri == 0
                    r.font.color.rgb = RGBColor.from_string('FFFFFF' if ri == 0 else '000000')
    d.add_paragraph().paragraph_format.space_after = Pt(0)
    return t

def link(d, label, url, size=9):
    par = d.add_paragraph()
    hyper = OxmlElement('w:hyperlink')
    rel = d.part.relate_to(url, 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink', is_external=True)
    hyper.set(qn('r:id'), rel)
    run = OxmlElement('w:r'); pr = OxmlElement('w:rPr')
    c = OxmlElement('w:color'); c.set(qn('w:val'), '1F4E79'); pr.append(c)
    sz = OxmlElement('w:sz'); sz.set(qn('w:val'), str(int(size*2))); pr.append(sz)
    run.append(pr); text = OxmlElement('w:t'); text.text = label; run.append(text)
    hyper.append(run); par._p.append(hyper)
    par.paragraph_format.space_after = Pt(3)

objectives = [
    ['O1', 'Complete the five priority user stories', 'All Must acceptance checks pass by the E submission. No open critical defect.'],
    ['O2', 'Make joining and speaking practical', 'At least 9 of 10 setup trials finish within 2 minutes. At least 27 of 30 approved speaking trials become audible within 3 seconds.'],
    ['O3', 'Test whether listeners hear questions more easily', 'At least 75% positive feedback from at least 5 volunteer reviewers. Report counts and the small sample.'],
    ['O4', 'Keep microphone use under explicit control', 'Zero unauthorized transmissions in misuse tests. No saved audio or transcripts. An alternative speaking route remains available.'],
]

requirements = [
    dict(id='R1', epic='Session access', feature='Temporary browser session and QR or code join', story='As a student, I want to join a temporary room session without installing an app so that I can participate quickly.', persona='Student', priority='Must', acceptance='A valid code joins the correct session. An expired or ended session rejects joining. Denying microphone permission leaves the interface usable. At least 9 of 10 timed setup trials finish within 2 minutes.', objective='O1; O2; O4', status='Proposed - stakeholder validation pending', evidence='ClassMic in-class SMART note; prior ClassMic charter', sprint='1'),
    dict(id='R2', epic='Instructor control', feature='Speaking queue with approval, mute, removal and end controls', story='As an instructor, I want to choose and stop speakers so that I can manage classroom discussion.', persona='Instructor', priority='Must', acceptance='One request appears once in the queue. Only the approved student may transmit. Mute, removal, timeout and end session stop transmission. Concurrent requests never create two active speakers. All control and misuse checks pass.', objective='O1; O4', status='Proposed - stakeholder validation pending', evidence='ClassMic in-class SMART note; prior ClassMic charter', sprint='1 and 2'),
    dict(id='R3', epic='Classroom audio', feature='One approved phone transmits to a laptop and connected speaker', story='As a student, I want to speak through my phone after approval so that other people in the room can hear my question.', persona='Student speaker and listener', priority='Must', acceptance='Microphone notice and permission precede capture. Audio transmits only during approved push to talk and stops on release or revocation. At least 27 of 30 scripted trials become audible within 3 seconds of approval when the tester is ready to speak. Report median and 95th percentile one-way audio delay separately. The 95th percentile must not exceed 150 milliseconds. No sustained feedback is acceptable.', objective='O1; O2; O3; O4', status='Proposed - device and AV validation pending', evidence='ClassMic SMART note; prior charter; technical feasibility still untested', sprint='2'),
    dict(id='R4', epic='Access and recovery', feature='Clear permission, connection and alternative participation paths', story='As a participant who cannot use a phone microphone, I want another way to contribute so that I am not excluded.', persona='Student with device or access barrier', priority='Must', acceptance='Permission denial and disconnection show a clear next step. The instructor can repeat a question or provide an available physical microphone. Keyboard navigation and readable control labels pass the agreed review. Rejoining never restores a previous speaking approval.', objective='O1; O4', status='Proposed - accessibility feedback pending', evidence='Prior ClassMic charter and plan; user validation pending', sprint='1 and 2'),
    dict(id='R5', epic='Evaluation', feature='Test record and anonymous feedback without stored speech', story='As the project manager, I want a record of test outcomes and feedback so that I can recommend whether further testing is worthwhile.', persona='Project manager', priority='Must', acceptance='The report covers every Must check, 10 setup trials, 30 speaking trials and feedback from at least 5 volunteers. At least 75% report easier audibility. Save results and minimal test metadata only. Report failures, sample limits and open defects. No audio or transcript is saved.', objective='O1; O2; O3; O4', status='Proposed - evaluation protocol pending', evidence='ClassMic SMART note and prior charter; sample sizes proposed for course feasibility', sprint='2'),
]

schedule = [
    ['S1', 'A2, B2 and charter review', '2026-09-10', '2026-09-12', '3', '', 'Both', 'Yes'],
    ['S2', 'Requirements and stakeholder input', '2026-09-13', '2026-09-15', '3', 'S1', 'Nurtore', 'Yes'],
    ['S3', 'Technical approach and device access', '2026-09-13', '2026-09-15', '3', 'S1', 'Temiko', 'Yes'],
    ['S4', 'Integrated D plan and review', '2026-09-16', '2026-09-19', '4', 'S2;S3', 'Both', 'Yes'],
    ['S5', 'Sprint 1 session and control workflow', '2026-09-20', '2026-09-26', '7', 'S4', 'Both', 'Yes'],
    ['S6', 'Sprint 2 audio integration and evaluation', '2026-09-27', '2026-10-03', '7', 'S5', 'Both', 'Yes'],
    ['S7', 'Closeout and final presentation', '2026-10-04', '2026-10-09', '6', 'S6', 'Both', 'Yes'],
]

effort = [
    ['Initiation and D planning', 18, 22, 40],
    ['Sprint 1', 18, 6, 24],
    ['Sprint 2', 14, 10, 24],
    ['Closing and presentation', 4, 8, 12],
    ['Total hours', 54, 46, 100],
]

# A2
d = base('ClassMic Aspiration', 'A2')
title(d, 'ClassMic Aspiration', 'A2   Revised project opportunity')
p(d, 'Our aim is to make student contributions easier for the whole classroom to hear while the instructor keeps control of the discussion. We propose a small ClassMic prototype to test whether phones can support that aim through the room speakers.')
h(d, 'Problem and opportunity')
p(d, 'A question from one student can help the whole class only if others hear it. When a voice does not carry across the room, the instructor may need to repeat the question or wait while a microphone reaches the student. Those workarounds can interrupt discussion. The size and frequency of this problem in our classes still need to be measured.')
p(d, 'ClassMic addresses the CONNECT theme through shared learning and direct classroom participation. The opportunity is to improve how people hear and respond to one another during an existing class.')
h(d, 'What we know')
p(d, 'The in-class SMART note proposes an instructor-controlled phone microphone, existing Wi-Fi and speakers, one active microphone, low delay and 75% positive feedback. It establishes the idea and intended outcome. It does not establish demand, tested performance or approval for a funded campus pilot.')
p(d, 'A commercial example, Biamp Crowd Mics, already supports phone microphones and moderator control [1]. This makes it a useful alternative to investigate. ClassMic must demonstrate the value of its proposed browser workflow before the team recommends a larger investment.')
h(d, 'Value proposition')
p(d, 'For students and instructors who find peer contributions hard to hear, ClassMic would provide a temporary speaking channel from a student phone to a room speaker. Students request a turn. The instructor approves the speaker and can stop the audio. The prototype will test whether this process is easier than the current classroom workaround.')
h(d, 'Initial project boundary')
p(d, 'The course project covers requirements, a focused prototype, controlled tests and a recommendation. Development uses two weekly sprints. A live classroom rollout, paid infrastructure and a six-week institutional pilot require a later decision.')

page(d, 'Major stakeholders')
table(d, ['Stakeholder', 'Need or decision'], [
    ['Students', 'Hear peer contributions and choose whether to use a phone microphone. Provide feedback on access and clarity.'],
    ['Instructors', 'Manage turns without distracting from teaching. Validate approval, mute and end controls.'],
    ['Project team', 'Temiko leads the prototype. Nurtore coordinates requirements and delivery. These are the selected planning roles.'],
    ['Course instructor', 'Review the revised idea, charter and course deliverables. Formal course acceptance remains open.'],
    ['IT AV and accessibility advisers', 'Review device, network, sound and access issues if a real room test becomes available. Participation is unconfirmed.'],
], [1.55, 5.55])
h(d, 'Initial success measures')
p(d, 'These are proposed course evaluation targets, not results or approved institutional service levels.')
table(d, ['Measure', 'Proposed target'], [
    ['Working scope', 'Complete five priority user stories and all Must acceptance checks.'],
    ['Setup and speaking', '9 of 10 setup trials within 2 minutes. 27 of 30 approved speaking trials audible within 3 seconds.'],
    ['Perceived audibility', 'At least 75% positive feedback from at least 5 volunteer reviewers. Report the actual numerator and denominator.'],
    ['Control and privacy', 'No unauthorized transmission in misuse tests and no saved audio or transcripts.'],
], [1.55, 5.55])
h(d, 'Thirty second executive pitch')
p(d, 'ClassMic is a classroom microphone concept that lets students request a turn from their phones and speak through the room speakers after instructor approval. We want to test whether it makes student questions easier to hear without adding classroom friction. Our course project will build and evaluate a small prototype in two weekly sprints, then recommend whether a larger pilot is worthwhile.')
h(d, 'Items requiring confirmation')
p(d, 'Use the revised aspiration as the working direction. Gather local audibility examples and request the course review before development. Check device access before testing. The selected one-way audio-delay target is a 95th percentile of no more than 150 milliseconds. The earlier 14-week pilot remains a future proposal.')
link(d, '[1] Biamp Crowd Mics product overview, reviewed 10 September 2026', BIO)
d.save(OUT / 'A_Aspiration/ClassMic_A2_Aspiration.docx')

# B2
d = base('ClassMic Business Case', 'B2')
title(d, 'ClassMic Business Case', 'B2   Investment decision for the course prototype')
h(d, 'Recommendation')
p(d, 'GO for a limited course prototype using existing resources. Build the browser workflow to test the ClassMic idea and complete the course learning objectives. Keep any funded classroom pilot at REVISE until local need, audio performance, support effort and a commercial comparison are established.')
p(d, 'The decision requested is permission to investigate, prototype and evaluate. The preliminary plan allows 100 student hours and assumes no new cash spending. Use these as the planning baseline and check actual capacity before each sprint.')
h(d, 'Opportunity and expected benefit')
p(d, 'ClassMic aims to let the room hear student questions without waiting for a microphone to reach the speaker. Students join a temporary browser session, request a turn and transmit audio only after instructor approval. The instructor can mute a speaker or end the session.')
p(d, 'Expected value is easier audibility and less interruption during discussion. The course also gains a concrete way to study requirements, project control and iterative delivery. Improved grades, increased participation and institutional savings remain hypotheses. The present evidence cannot support a monetary return calculation.')
h(d, 'Evidence and its limits')
table(d, ['Evidence', 'What it supports', 'What remains unknown'], [
    ['In-class SMART note', 'The phone microphone concept and intended outcomes.', 'How often students miss questions and whether they want this workflow.'],
    ['Existing ClassMic drafts', 'A starting scope and list of risks.', 'Approval, available resources and measured feasibility.'],
    ['Biamp product information [1]', 'A commercial product offers the same core microphone and moderator workflow.', 'Local price, room fit and comparison with our prototype.'],
    ['Browser media specification [3]', 'Browser microphone capture is technically supported with permission controls.', 'Performance and compatibility on the devices and network we will use.'],
], [1.6, 2.6, 2.9], font=10)
h(d, 'If we do nothing')
p(d, 'Continue current classroom practice. This avoids development and maintenance work, but leaves any audibility problem unchanged. We should observe the current process before claiming that ClassMic improves it.')

page(d, 'Alternatives appraisal')
p(d, 'Compare the same classroom objective across three options. The judgments below are preliminary. Unknown costs and untested adoption are left unresolved rather than converted into a precise score.')
table(d, ['Criterion', 'Improve current process', 'Buy or adapt', 'Custom ClassMic prototype'], [
    ['Strategic fit', 'Support discussion by repeating questions or using an existing microphone.', 'Existing audience microphone workflow may fit the classroom.', 'Focused test of browser access and instructor control.'],
    ['Expected benefit', 'May solve the problem with a small change in teaching practice.', 'Phone speaking and moderation are established product capabilities [1].', 'Benefit must be demonstrated against the current workaround.'],
    ['Cost and value', 'No new platform. Staff time and equipment availability still matter.', 'Room hardware and installation require a quote [2]. Total cost unknown.', 'Proposed 100 student hours. No incremental cash assumed. Maintenance is unresolved.'],
    ['Feasibility', 'Lowest technical complexity if the existing workaround is adequate.', 'Needs procurement, AV integration and a product trial.', 'Browser, network, feedback and permission behavior need testing.'],
    ['Schedule', 'Can begin after agreement on the classroom routine.', 'Delivery and installation dates require a supplier response.', 'Two weekly sprints are planned. Audio is the main schedule risk.'],
    ['Security and privacy', 'No new microphone software or data flow.', 'Review actual configuration, data handling and support access.', 'Control data capture and avoid saved speech. The team owns testing.'],
    ['Operational support', 'Uses current teaching and AV support arrangements.', 'Supplier and institution divide support duties. Terms are unknown.', 'Two students support the prototype. Production support is outside scope.'],
    ['Stakeholder adoption', 'Familiar process, but repeated questions may remain disruptive.', 'Test the attendee and moderator workflow with real users.', 'No-install browser access is a hypothesis to validate with students.'],
], [1.05, 2.0, 2.0, 2.05], font=9.5)
h(d, 'Reason for the recommendation')
p(d, 'A custom prototype fits the course requirement to build and evaluate selected user stories. That educational reason does not establish that a university should build its own microphone service. A commercial trial and an improved current process remain credible options for a later institutional decision.')

page(d, 'Proposed delivery and resources')
table(d, ['Stage', 'Planned dates', 'Decision or output'], [
    ['Initiation and planning', '10–19 September', 'Review A2 B2 and C1. Complete requirements and D plan before the first sprint.'],
    ['Sprint 1', '20–26 September', 'Temporary sessions, permission flow, queue and instructor controls. Review the working flow.'],
    ['Sprint 2', '27 September–3 October', 'Integrate audio, test misuse and recovery, collect volunteer feedback and report results.'],
    ['Closing', '4–9 October', 'Assess delivery against the charter and prepare the final recommendation.'],
], [1.6, 1.7, 3.8])
p(d, 'The LMS deadlines are D on 22 September, E on 4 October and F on 10 October, each at 00:00 as displayed. Schedule D review for 19 September so that both full weekly sprints fit. The E upload area opens on 26 September; this controls uploading and does not state a restriction on earlier construction. If the D review is late, resolve the resulting schedule conflict before changing the sprint lengths.')
h(d, 'Preliminary cost and effort')
table(d, ['Work', 'Temiko hours', 'Nurtore hours', 'Total hours'], effort, [3.2, 1.3, 1.3, 1.3], center_cols=(1, 2, 3))
p(d, 'Cash plan: AED 0 of new spending, conditional on reusing available phones, laptops, speakers and permitted tools. Student effort is in kind, not free of opportunity cost. No labor rate or supplier quote is available, so this case does not invent a cash equivalent. Any paid hosting, equipment or account upgrade requires a revised estimate and approval before commitment.', lead='Cash plan:')
h(d, 'Feasibility conditions')
p(d, 'Confirm access to at least two suitable phones, a laptop and a speaker. Test microphone permission and browser behavior early. Use volunteer tests in a controlled setting. Room access, advisers and participants are not yet committed.')
p(d, 'The selected prototype target is a 95th percentile one-way audio delay of no more than 150 ms across 30 timing trials. This is a planning target, not a tested result. Also assess intelligibility and sustained feedback; a delay number alone does not establish usable room sound. A commercial comparison gives context [2].')

page(d, 'Decision gates and principal risks')
table(d, ['Gate', 'Evidence required', 'Decision'], [
    ['Before Sprint 1', 'Team and instructor review, five-story scope, resources, test plan and proposed schedule.', 'GO with the prototype only if the plan is feasible. Otherwise revise scope or dates.'],
    ['After Sprint 1', 'Session and queue demo, microphone feasibility check, feedback and updated backlog.', 'Continue core audio integration. Drop optional work if capacity or compatibility is weak.'],
    ['End of Sprint 2', 'All Must checks. 9/10 setups within 2 minutes, 27/30 speaking trials within 3 seconds, and at least 75% positive feedback from at least 5 reviewers. No critical defect.', 'Complete the course evaluation. Report unmet targets and test limitations.'],
    ['Any future live pilot', 'Local need, acceptable audio, named support owner, privacy and access review, supplier comparison and funded estimate.', 'REVISE until evidence is sufficient. Seek a separate authorization.'],
], [1.35, 3.45, 2.3], font=10)
h(d, 'Risks that could change the decision')
p(d, 'Feedback or network delay may make the audio unsuitable. Browser differences may prevent some students from using it. Weak instructor controls could permit unwanted transmission. Limited student time or unavailable reviewers could prevent a useful evaluation. Stop an affected test if control or privacy fails, preserve a physical or instructor-repeat alternative, and narrow the prototype before adding resources.')
h(d, 'Investment conclusion')
p(d, 'The present evidence supports a bounded learning project. It does not justify a funded institutional rollout or a claim that custom development is the best long-term purchase decision. At closeout, recommend further testing only if the prototype demonstrates useful audibility, workable instructor control and a realistic support path.')
h(d, 'Items requiring confirmation')
p(d, 'Confirm the revised A2 and B2, team capacity, early D approval, test equipment and volunteers. Use the selected 150 ms delay target and the supplied test-data protocol. Record the instructor’s decision. Do not infer approval from an LMS upload.')
h(d, 'Sources')
link(d, '[1] Biamp Crowd Mics product overview', BIO)
link(d, '[2] Biamp Crowd Mics FAQ, including quote requirements and audio latency', FAQ)
link(d, '[3] W3C Media Capture and Streams specification', W3)
p(d, 'External sources reviewed 10 September 2026. Local inputs: the ClassMic in-class SMART note and existing charter and plan. Course requirements: the LMS B, C, D, E and F pages and D Phase Labs attachment. No stakeholder interviews, supplier quotation or pilot results were supplied.', size=9)
d.save(OUT / 'B_Business_Case/ClassMic_B2_Business_Case.docx')

# C1
d = base('ClassMic Project Charter', 'C1 draft for review', body_size=10.5)
title(d, 'ClassMic Project Charter', 'C1   Develop the Project Charter')
p(d, 'Draft authorization for the course prototype. This charter carries forward the revised ClassMic A2 and B2. Team and instructor approval remain to be recorded. All targets and resource estimates below are proposed.', lead='Draft authorization for the course prototype.', size=10)
h(d, '1 Purpose and justification')
p(d, 'Test whether student phones can make classroom questions easier to hear while the instructor controls speaking turns. B2 recommends a small course prototype using existing resources. Comparing it with current practice will inform a later pilot decision. The project supports the CONNECT theme through shared classroom learning.')
h(d, '2 Measurable objectives and success criteria')
table(d, ['ID', 'Target by the E submission'], [
    ['O1', 'Complete five priority user stories. Pass all Must acceptance checks with no open critical defect.'],
    ['O2', 'At least 9 of 10 setup trials within 2 minutes and 27 of 30 approved speaking trials audible within 3 seconds.'],
    ['O3', 'At least 75% positive audibility feedback from at least 5 volunteer reviewers. Report counts and sample limits.'],
    ['O4', 'Zero unauthorized transmissions in misuse tests. No saved audio or transcripts. Preserve an alternative speaking route.'],
], [.45, 6.65], font=10)
p(d, 'Use scripted trials with a ready speaker. Measure audio transport delay separately and report its median and 95th percentile. The selected prototype target is a 95th percentile one-way delay of no more than 150 milliseconds. The three-second target includes speaking activation, not just network delay.', size=9.5)
h(d, '3 High level scope and exclusions')
p(d, 'Included: temporary browser sessions, QR or code join, microphone permission, speaking queue, instructor approval and mute/end controls, one active push-to-talk phone, speaker output, recovery guidance and controlled evaluation.', lead='Included:')
p(d, 'Excluded: live teaching deployment, a six-week institutional pilot, recording or transcripts, permanent profiles, speech-based grading or attendance, simultaneous microphones, remote learning and paid production infrastructure.', lead='Excluded:')
h(d, '4 Deliverables and milestones')
table(d, ['Planned date', 'Deliverable'], [
    ['By 19 September', 'Reviewed A2, B2 and charter, PRD, scope and integrated D plan.'],
    ['20–26 September', 'Sprint 1 demo of session access and instructor controls.'],
    ['27 September–3 October', 'Sprint 2 audio prototype, test record, feedback and evaluation.'],
    ['4–9 October', 'Closeout review and final presentation.'],
], [1.45, 5.65], font=10)
p(d, 'LMS cutoffs: D 22 September, E 4 October, F 10 October, all at 00:00 as displayed. The plan schedules D review by 19 September. The E upload opening date does not state a restriction on starting work.', size=9.5)

page(d, '5 High level requirements')
p(d, 'The browser requests explicit permission before microphone capture. Only an instructor-approved participant can transmit. Releasing push to talk, muting, disconnecting or ending the session must stop transmission. Expired sessions reject access. The prototype saves no audio or transcripts. Keep anonymous test counts, timings and defect records until 30 days after course feedback, then delete participant-level records. Participants who cannot use the phone workflow can ask the instructor to repeat their question or use an available physical microphone.')
h(d, '6 Initial risks assumptions and constraints')
p(d, 'Risks: feedback, network delay, device incompatibility, unwanted transmission, insufficient testing and limited team capacity. Test the audio path early, retain instructor stop controls, record defects and remove optional work if the schedule slips. Stop affected tests for suspected privacy or control failures.', lead='Risks:')
p(d, 'Assumptions: suitable phones, a laptop and speaker are available; volunteers can review the prototype; the team can provide about 100 hours in total; and the course instructor accepts the proposed sequence. Equipment access, reviewer availability and team capacity require confirmation.', lead='Assumptions:')
p(d, 'Constraints: two weekly construction sprints, five priority stories, the published course deadlines and no new spending without approval. The initial cash forecast is AED 0 because existing resources are assumed. Live classroom use requires a separate authorization.', lead='Constraints:')
h(d, '7 Sponsor and major stakeholders')
p(d, 'Course acceptance authority: the BUS 2010 instructor. Formal charter approval and the named sponsor record remain open. Stakeholders are students as speakers and listeners, instructors as operators, and the two-person project team. IT, AV, privacy and accessibility advisers may support reviews if available. No university department or external sponsor is committed by this draft.')
h(d, '8 Project manager and authority')
p(d, 'Planning roles: Nurtore Arynuruly is project manager. Temiko Machavariani is technical lead. On approval, the project manager coordinates work, tracks risks and changes, and accepts internal work against the PRD. Temiko leads implementation and technical testing. The course instructor accepts assessed deliverables and material changes to scope or timing. Neither team member may commit university funds or authorize live deployment.')
h(d, 'Items requiring confirmation')
p(d, '1. Record team and instructor approval of the revised A2, B2, C1 and role assignments. Confirm that D may use this approved C1 as the charter input labelled C2 in its brief.', size=10)
p(d, '2. Confirm formal D review by 19 September, equipment access and volunteer availability. The plan defines the test protocol, delay target and data handling.', size=10)
p(d, '3. Add genuine team discussion images and any missing individual input. Temiko’s supplied SMART note and an AI-assisted brainstorming proposal accompany this charter; neither claims a completed group meeting.', size=10)
p(d, 'Approval record   Decision and conditions: ______________________________', size=10)
p(d, 'Course instructor: _____________________   Date: __________', size=10)
p(d, 'Project manager: ______________________   Date: __________', size=10)
d.save(OUT / 'C_Charter/ClassMic_C1_Project_Charter.docx')

# D1 to D7
d = base('ClassMic Integrated Project Plan', 'D planning draft')
title(d, 'ClassMic Integrated Project Plan', 'D1 to D7   Course prototype and two weekly sprints')
h(d, 'D1 Development approach')
p(d, 'Use a hybrid approach. Complete the D planning outputs in sequence, then build and review the prototype in two weekly E sprints. Close the course project in F using the actual results. The revised A2, B2 and C1 define the proposed scope. This plan is ready for review, but it does not claim stakeholder validation or charter approval.')
table(d, ['Predictive planning', 'Adaptive construction'], [
    ['Agree on the classroom problem, five-story scope, responsibilities and acceptance boundary before development.', 'Use Sprint 1 to test the session and instructor workflow, then revise priorities from feedback.'],
    ['Set the course milestones, effort forecast, risk responses and test approach. Confirm permission and data rules.', 'Use Sprint 2 for the audio path, recovery, control tests and evaluation. Correct defects before optional polish.'],
    ['Record approvals and changes. Keep the D baseline stable enough to track progress.', 'Update the Notion PRD and backlog after each review while keeping the charter objectives traceable.'],
], [3.55, 3.55])
h(d, 'Why the approach fits')
p(d, 'The assessed outputs and deadlines are fixed, so the team needs a plan before building. The audio path, browser behavior and instructor workflow are uncertain, so short demonstrations and feedback should guide implementation. Two students can coordinate through a small backlog without creating separate management roles for every Scrum responsibility.')
h(d, 'Controls and assumptions')
p(d, 'No live classroom rollout is included. Controlled tests need willing participants, available equipment and an agreed privacy approach. The course instructor reviews material scope or schedule changes. Any university room test or new spending requires the relevant permission before it takes place.')
p(d, 'The current brief names D7 in its overview but supplies detailed prompts only for D1–D6. The D7 review in this document checks consistency across those outputs. D1 remains unvalidated until the team and instructor record their review.')
h(d, 'Decision to record')
p(d, 'The selected baseline is five stories, 100 student hours and the two listed sprint dates. Formal course review and actual resource availability remain open. The older 14-week, three-room pilot is a possible later project and is outside this plan.')

page(d, 'D2 Scope statement')
p(d, 'ClassMic will test whether a student phone can carry a question through a speaker after instructor approval. The project delivers a PRD and backlog, a browser prototype covering five user stories, controlled test evidence, volunteer feedback and a recommendation. It includes temporary session access, microphone permission, speaking requests, one approved push-to-talk microphone, instructor stop controls and a recovery path.')
p(d, 'The team will not deliver a production service, run a six-week live classroom pilot, save audio or transcripts, identify speakers by voice, grade participation or support simultaneous microphones. It assumes access to existing devices and a speaker, approximately 100 student hours and willing reviewers. Completion requires all Must acceptance checks, the charter’s proposed evaluation measures, an honest account of failures and no open critical control defect. Course deadlines and two weekly construction sprints constrain delivery.')
h(d, 'Requirements summary')
table(d, ['ID', 'Requirement', 'Charter link'], [
    [r['id'], r['feature'], r['objective']] for r in requirements
], [.45, 5.4, 1.25], font=10)
h(d, 'PRD structure and ownership')
p(d, 'Nurtore maintains the PRD. Temiko updates implementation status and test evidence. The supplied Notion import table uses the required fields: Epic, Feature/Requirement, User Story, Persona, Priority, Acceptance Criteria, Charter Objective and Status. The separate PRD text contains the context, personas, scope and open questions.')
p(d, 'Until the team imports the files and provides the page link, the Notion requirement remains open. After import, use that page as the maintained requirements record and retain exports at review milestones.')
h(d, 'Evidence boundary')
p(d, 'R1–R3 follow the in-class ClassMic concept and prior charter. R4 and R5 carry forward access and evaluation needs from the prior drafts. Each remains a proposed requirement until students and an instructor review it. No interview, survey or usability result has been invented.')

page(d, 'D2 Acceptance and priorities')
table(d, ['Story', 'Acceptance boundary', 'Iteration'], [
    ['US1 / R1', 'Join the intended session. Reject expired access. Handle denied permission. Meet the setup target.', 'Sprint 1'],
    ['US2 / R2', 'Queue once, approve one speaker, and stop audio through mute, removal, timeout and end. Pass misuse tests.', 'Both'],
    ['US3 / R3', 'Transmit only with permission, approval and push to talk. Meet the speaking-start target. Measure audio delay separately.', 'Sprint 2'],
    ['US4 / R4', 'Explain denied permission and connection loss. Preserve an alternative speaking path and clear controls.', 'Both'],
    ['US5 / R5', 'Record all acceptance results and actual feedback counts. Save no speech. Report defects and limits.', 'Sprint 2'],
], [1.1, 4.9, 1.1], font=10)
h(d, 'Priority decisions')
p(d, 'Must: R1–R5 form the minimum complete prototype and evaluation. A failed core control, audio path or evidence requirement cannot be hidden by marking the prototype complete.')
p(d, 'Should: simple connection checks and a short operator guide, where they help reviewers use the prototype. These refine the existing stories and must not delay acceptance tests.')
p(d, 'Could: visual polish after the core workflow passes. Any new capability needs a requirement and evidence before it enters a sprint.')
p(d, 'Will not have: recording, transcript storage, voice identification, attendance, grading, persistent profiles, captions, polling, remote participation or campus deployment in this course release.')
h(d, 'Open questions for validation')
p(d, 'Which classroom situations cause audibility problems, and how often? Will instructors operate a queue during discussion? Which devices and network can the team test? What delay is acceptable with the available speaker? Which volunteer reviewers can participate? What minimum metadata, access and deletion rules will the team use?')
h(d, 'Traceability and change control')
p(d, 'Every backlog entry carries R1–R5 and O1–O4 references. Each test result names its requirement, method, device and outcome. Nurtore records requested changes and their effect on the scope, effort and course deadline. The instructor reviews material changes before the team treats them as approved.')

page(d, 'D3 Work breakdown structure')
p(d, 'The WBS contains 18 elements, including six summary elements. Each leaf has an owner and an identifiable output. Requirements R1–R5 map to the five user stories in the PRD.')
wbs = [
    ['1.0', 'Project authorization and control', 'Summary', 'Nurtore'],
    ['1.1', 'Reviewed A2, B2 and charter', 'Decision and confirmation record', 'Nurtore'],
    ['1.2', 'Status and change records', 'Current actions, risks and decisions', 'Nurtore'],
    ['2.0', 'Requirements and plan', 'Summary', 'Nurtore'],
    ['2.1', 'Stakeholder input and PRD', 'Five validated or explicitly open stories', 'Nurtore'],
    ['2.2', 'Integrated D plan', 'Scope, schedule, risk, quality, cost and team plan', 'Nurtore'],
    ['3.0', 'Sprint 1 increment', 'Summary', 'Temiko'],
    ['3.1', 'Session access and permissions', 'R1 working flow', 'Temiko'],
    ['3.2', 'Instructor queue and controls', 'R2 control flow and review', 'Temiko'],
    ['4.0', 'Sprint 2 increment', 'Summary', 'Temiko'],
    ['4.1', 'Audio path and recovery', 'R3 and R4 integrated prototype', 'Temiko'],
    ['4.2', 'Acceptance and evaluation', 'R1–R5 test record and feedback', 'Nurtore'],
    ['5.0', 'Delivery evidence', 'Summary', 'Nurtore'],
    ['5.1', 'E demonstration', 'Working demo and defect record', 'Temiko'],
    ['5.2', 'Updated requirements record', 'PRD with links to tests and decisions', 'Nurtore'],
    ['6.0', 'Closing', 'Summary', 'Nurtore'],
    ['6.1', 'Outcome and lessons review', 'Actual results against O1–O4', 'Nurtore'],
    ['6.2', 'Final presentation and handover', 'F presentation and retained project files', 'Nurtore'],
]
table(d, ['WBS', 'Element', 'Output', 'Owner'], wbs, [.5, 2.4, 3.15, 1.05], font=9.5)
p(d, 'Required human input: the team must supply its own hand-drawn or Post-it WBS and dependency network. This typed WBS is a planning proposal and does not replace those images.', size=10)

page(d, 'D3 Schedule and release plan')
p(d, 'All dates below are proposed. Durations are calendar-day windows, not full-time labor. The plan completes D early so that two full weekly sprints fit before E is due. D needs review by 19 September. The E upload area opens on 26 September; its opening date does not prohibit earlier work.')
table(d, ['ID and activity', 'Dates', 'Days', 'Depends on'], [
    [s[0] + '  ' + s[1], date.fromisoformat(s[2]).strftime('%d %b') + '–' + date.fromisoformat(s[3]).strftime('%d %b'), s[4], s[5] or 'None'] for s in schedule
], [3.5, 1.6, .5, 1.5], font=10, center_cols=(2,))
h(d, 'Gantt view')
gantt = table(d, ['Activity', '10–12 Sep', '13–19 Sep', '20–26 Sep', '27 Sep–3 Oct', '4–9 Oct'], [
    ['A2, B2 and C1 review', 'Work', '', '', '', ''],
    ['Requirements and D review', '', 'Work', '', '', ''],
    ['Sprint 1', '', '', 'Work', '', ''],
    ['Sprint 2 and E', '', '', '', 'Work', ''],
    ['Closing and F', '', '', '', '', 'Work'],
], [2.3, .96, .96, .96, .96, .96], font=9, center_cols=(1,2,3,4,5))
for ri, row in enumerate(gantt.rows[1:], 1):
    for cell in row.cells[1:]:
        if cell.text == 'Work':
            cell._tc.get_or_add_tcPr().find(qn('w:shd')).set(qn('w:fill'), 'BED3E1')
h(d, 'Critical dependencies and resource conflicts')
p(d, 'S1 precedes parallel S2 and S3. Both must finish before S4, then S5, S6 and S7 follow in order. These are controlling paths with no planned buffer before the internal milestone dates. E is due at 00:00 on 4 October, so finish and check the upload on 3 October. F follows the same pattern on 9 October.')
p(d, 'Audio feasibility and Temiko’s availability have low-confidence estimates. Nurtore’s testing and evidence work overlap Sprint 2. Reserve review time in each sprint. If D approval slips past 19 September, ask the instructor to resolve the two-week schedule conflict. Do not silently shorten either sprint or assume an extension.')

page(d, 'D4 Risk register')
p(d, 'The register has six threats followed by two opportunities. Probability and impact are qualitative estimates awaiting team review. High impact means a core objective, deadline or control could fail. Medium impact means rework or weaker evidence. Low impact means a local inconvenience.')
risks = [
    ['T1', 'Speaker feedback or delay makes questions hard to understand.', 'High', 'High', 'Test the audio path early at low volume. Adjust the setup and retain a fallback.', 'Temiko'],
    ['T2', 'Browser or network differences prevent some devices from joining or transmitting.', 'Medium', 'High', 'Test available devices and permissions early. Publish the tested combination.', 'Temiko'],
    ['T3', 'A control or configuration defect permits unwanted capture or transmission.', 'Medium', 'High', 'Keep one approved speaker. Test denial, expiry, mute and end. Save no speech.', 'Temiko'],
    ['T4', 'Competing coursework or late D approval leaves too little time for two sprints.', 'High', 'High', 'Confirm capacity and the early review date. Reserve testing time and drop optional work.', 'Nurtore'],
    ['T5', 'Unavailable or excluded reviewers leave too little useful feedback.', 'Medium', 'Medium', 'Invite willing reviewers early. Provide an alternative speaking path and report sample limits.', 'Nurtore'],
    ['T6', 'Unclear audio expectations or extra features expand the work after planning.', 'Medium', 'High', 'Agree on acceptance before building. Assess changes against the five-story boundary.', 'Nurtore'],
    ['OP1', 'Existing equipment or adviser support becomes available for a better sound test.', 'Medium', 'Medium', 'Seek access and a short review without making it a delivery dependency.', 'Nurtore'],
    ['OP2', 'Standard browser capabilities work well enough to reduce integration effort.', 'Medium', 'Medium', 'Prototype the transport early and reuse supported components after review.', 'Temiko'],
]
table(d, ['ID', 'Event', 'Prob', 'Impact', 'Response', 'Owner'], risks, [.4, 2.2, .65, .65, 2.25, .95], font=9.5)
h(d, 'Risk ownership')
p(d, 'Owners review triggers during the team check-in and at each sprint review. Record a risk as an issue when it occurs. Nurtore escalates a threatened course deadline or scope change to the instructor. Temiko stops an affected test immediately if microphone control fails.')

page(d, 'D4 Triggers and contingency actions')
table(d, ['ID', 'Observable trigger', 'Contingency or opportunity action'], [
    ['T1', 'Sustained feedback or listeners cannot follow the test speech.', 'Stop audio. Retest at safe settings or use another speaker setup. Report an unmet audio requirement if unresolved.'],
    ['T2', 'A required test device cannot join or transmit after troubleshooting.', 'Use the documented supported device for the demo and record the compatibility failure. Do not claim wider support.'],
    ['T3', 'Any unapproved transmission or saved speech appears.', 'Stop the test and disable the affected path. Correct the defect and rerun misuse tests before resuming.'],
    ['T4', 'D lacks approval by 19 September or a core sprint item exceeds available effort.', 'Escalate dates or scope for instructor review. Remove optional polish. Preserve test and evidence time.'],
    ['T5', 'Fewer than five willing reviewers are available or a reviewer lacks an equivalent participation route.', 'Arrange another controlled review or report insufficient evidence. Never invent responses or force participation.'],
    ['T6', 'A request adds a new feature or changes the agreed acceptance threshold.', 'Log impact and obtain a decision before adding work. Keep the original result visible if a target changes.'],
    ['OP1', 'An adviser or existing room setup is available without new spending.', 'Schedule an additional comparison test with permission. Record the configuration and findings.'],
    ['OP2', 'Early transport checks pass with less effort than estimated.', 'Use the released time for device and accessibility testing. Do not expand the release boundary automatically.'],
], [.55, 2.55, 4.0], font=10)
h(d, 'Response records')
p(d, 'For each triggered event, record the date, evidence, owner, action and effect on the next milestone. Close it only when the response has been checked. Opportunities count as benefits only when they occur and the team records what changed.')

page(d, 'D4 Quality and acceptance plan')
p(d, 'Quality assurance prevents avoidable defects: review requirements before coding, check changes against the PRD, review microphone data flow and keep an agreed definition of done. Inspection and testing check the resulting prototype: run functional, permission, misuse, audio and user-review scenarios and record their outcomes.')
table(d, ['Test area', 'Acceptance evidence', 'Owner'], [
    ['Scope and controls', 'R1–R5 checklist. All Must checks pass. Tests cover two simultaneous requests, mute, removal, expiry, disconnect and end.', 'Both'],
    ['Setup and speaking start', '10 setup trials, at least 9 within 2 minutes. 30 ready-speaker trials, at least 27 audible within 3 seconds of approval.', 'Nurtore'],
    ['Audio quality and delay', 'Measure 30 one-way delay trials. Report median and 95th percentile; the latter must be at most 150 ms. Test intelligibility and sustained feedback. Record any failure.', 'Temiko'],
    ['Privacy and data', 'Configuration and storage review shows no saved speech. Misuse tests show no unauthorized transmission. Save minimal test metadata only.', 'Temiko'],
    ['Access and recovery', 'Permission denial, reconnect and alternative participation checks. Review readable labels and keyboard access.', 'Both'],
    ['User feedback', 'At least 5 willing reviewers compare audibility with the current workaround. At least 75% positive, with counts, response wording and limitations reported.', 'Nurtore'],
], [1.4, 4.7, 1.0], font=10)
h(d, 'Measurement protocol')
p(d, 'Start setup timing when the instructor opens a new session and end when a test student is ready to request a turn. Use a scripted ready speaker for activation timing. For transport delay, use an agreed synthetic test signal and one calibrated timing method, not unsynchronized phone clocks. Retain timing results only. Note the device, browser and network for every trial.')
p(d, 'Ask reviewers whether questions were easier to hear than through the baseline workaround. Use the same test setting and explain the small convenience sample. With five responses, four positive responses are 80% and meet the proposed 75% target. This does not establish a campus-wide effect.')
h(d, 'Defects and acceptance')
p(d, 'Critical defects include unwanted transmission, saved speech or a failed stop control. Stop affected testing and require a verified fix. Other defects need an owner, priority and retest result. A failed Must check remains open in the report. The team reviews the evidence and the instructor accepts the assessed deliverable. Record actual outcomes during E.')

page(d, 'D5 Cost and resource plan')
p(d, 'The course prototype uses a provisional estimate of 100 student hours. These are planning allowances, not recorded time. Nurtore tracks actual effort and the forecast after each sprint. No wage rates or supplier quotations were supplied.')
table(d, ['Work', 'Temiko hours', 'Nurtore hours', 'Total hours'], effort, [3.2, 1.3, 1.3, 1.3], center_cols=(1,2,3))
table(d, ['Resource', 'Planning basis', 'Confirmation needed'], [
    ['Phones, laptop and speaker', 'Reuse available equipment. Incremental cash allowance AED 0.', 'Availability and compatibility. Any purchase changes the estimate.'],
    ['Development and hosting', 'Use already available tools and an approved local or existing environment.', 'Permitted environment and access. No paid subscription assumed.'],
    ['Notion PRD', 'Prepare importable PRD content and the required database fields.', 'Existing team account or page and access preferences.'],
    ['Volunteer and adviser time', 'Participation by agreement. No cash allowance or external commitment.', 'Willing reviewers and any technical advice.'],
], [1.5, 3.3, 2.3], font=10)
h(d, 'Cash baseline and uncertainty')
p(d, 'Provisional new cash spending is AED 0 if all assumed resources are available. This is a reuse condition, not evidence that a production system costs nothing. Any paid requirement must have a quote, an owner and approval before commitment. The previous AED 60,000 pilot ceiling does not apply to this course prototype.')
p(d, 'The effort estimate has low confidence until the early audio test. Temiko’s 18-hour Sprint 1 allowance and 14-hour Sprint 2 allowance are the main capacity limits. Reserve Nurtore’s evaluation time rather than shifting all testing to the final evening.')
h(d, 'Lower resource scenario')
p(d, 'If total effort falls to about 70 hours, first remove polish and extra device coverage. Keep the permission, one-speaker, stop-control and no-recording requirements. Seek approval for a reduced evaluation scope if necessary. A click-through mockup can still explain the workflow, but it cannot count as passing live-audio acceptance. Record that change instead of claiming the original scope was delivered.')

page(d, 'D6 Team communications and engagement')
table(d, ['Role', 'Proposed responsibility and authority'], [
    ['Nurtore Arynuruly', 'Project manager and requirements lead. Maintain PRD, schedule, risk and decision records. Coordinate reviews and evaluation.'],
    ['Temiko Machavariani', 'Technical lead. Implement and test the prototype. Maintain technical evidence and stop tests for control failures.'],
    ['Course instructor', 'Proposed course sponsor and acceptance authority. Review material scope and schedule changes. Confirm name and role.'],
    ['Student and instructor reviewers', 'Volunteer feedback on audibility, access and controls. No participation is assumed.'],
], [1.65, 5.45])
h(d, 'Working agreements')
p(d, 'Keep requirements in the agreed Notion PRD after import. Name one owner for each active task and link its evidence. Review changes that affect microphone behavior before merging or demonstrating them. Record meeting decisions and unresolved disagreements. Do not attribute AI-written material or planned feedback to a person as a completed contribution.')
table(d, ['Audience and purpose', 'Proposed cadence', 'Channel and owner'], [
    ['Team actions, blockers and risks', 'Ten-minute check-in on workdays; a 30-minute review at each sprint boundary.', 'Telegram for daily coordination; Notion for decisions. Nurtore keeps the action log.'],
    ['Sprint planning and review', 'At the start and end of each weekly sprint.', 'Working demo and PRD update. Both attend.'],
    ['Instructor decisions', 'Before Sprint 1 and when scope or deadlines are threatened.', 'LMS messaging or the next class review. Nurtore prepares the decision request.'],
    ['Volunteer feedback', 'After consent and at the agreed test session.', 'Use the supplied anonymous review form after a controlled demonstration. Nurtore records results.'],
    ['Technical or access advice', 'Before a room test, if an adviser is available.', 'Agreed review route. Temiko records the recommendation.'],
], [2.2, 2.35, 2.55], font=10)
h(d, 'Escalation and engagement')
p(d, 'Stop affected testing immediately for unwanted transmission or saved speech. Raise an unresolved Must defect, unavailable resource or threatened milestone at the next team check-in and promptly seek the instructor’s decision when it changes scope or timing. Ask priority reviewers how and when they prefer to participate. Collect no private contact details in the public project record.')
p(d, 'The companion brainstorming supplement explains the selected communication plan. It is AI-assisted planning prepared now. The D6 request for original pre-AI human notes remains open; this supplement does not claim to recreate those notes.', size=10)

page(d, 'D7 Integrated plan review')
p(d, 'This review links the planning outputs and identifies the remaining decisions. It does not mark the project approved or the prototype tested.')
table(d, ['Check', 'Current position', 'Next action'], [
    ['One project across A–D', 'All revised documents use ClassMic and the two-sprint course prototype.', 'Use the course prototype as the working scope; obtain formal course acceptance.'],
    ['Business case to charter', 'GO applies to a bounded prototype. A funded live pilot needs another decision.', 'Record A2 B2 and C1 review.'],
    ['Scope to backlog', 'Five Must stories map to O1–O4. Detailed acceptance fields are ready for Notion.', 'Import PRD and record the page link.'],
    ['Backlog to WBS', '18 WBS elements cover planning, two increments, evidence and closing.', 'Add genuine WBS and dependency images.'],
    ['Schedule to capacity', 'Two seven-day sprints require D review by 19 September. Effort totals 100 hours.', 'Complete D review by 19 September and check available effort at sprint planning.'],
    ['Risks to quality', 'Six threats and two opportunities have triggers, responses and owners.', 'Track the risk triggers and apply the selected 150 ms delay target.'],
    ['Resources to cash', 'AED 0 incremental cash assumes reuse. There are no quotes or paid commitments.', 'Verify access or revise the estimate before spending.'],
    ['Team to evidence', 'Responsibilities and communication proposals are documented.', 'Supply pre-AI communication notes and charter discussion evidence.'],
], [1.4, 3.25, 2.45], font=10)
h(d, 'Approval record to complete')
p(d, 'Review decision and conditions: ____________________________________')
p(d, 'Course instructor or agreed approver: __________________   Date: ________')
p(d, 'Project manager: __________________________________   Date: ________')
h(d, 'Requirements used for this revision')
p(d, 'The current D Phase Labs attachment controls the D1 one-page limit, D2 two-page limit outside Notion, two weekly sprints, no more than 20 WBS elements, and no more than eight risk events including two opportunities. These replace the older generic planning quantities. The attachment refers to C2 even though the visible charter assignment requests C1; instructor confirmation is recorded as an open item.')
link(d, 'D assignment and current attachment', 'https://lms.mbzuai.ac.ae/mod/assign/view.php?id=38962')
link(d, 'E execution assignment', 'https://lms.mbzuai.ac.ae/mod/assign/view.php?id=38965')
link(d, 'F closing assignment', 'https://lms.mbzuai.ac.ae/mod/assign/view.php?id=38968')
d.save(OUT / 'D_Develop_Plans/ClassMic_D_Integrated_Plan.docx')

# Notion-ready requirements and schedule data
fields = ['Epic', 'Feature/Requirement', 'User Story', 'Persona', 'Priority', 'Acceptance Criteria', 'Charter Objective', 'Status']
with (OUT / 'D_Develop_Plans/ClassMic_Notion_Requirements.csv').open('w', newline='', encoding='utf-8-sig') as f:
    writer = csv.writer(f); writer.writerow(fields)
    for r in requirements:
        writer.writerow([r['epic'], r['id'] + ' ' + r['feature'], r['story'], r['persona'], r['priority'], r['acceptance'], r['objective'], r['status']])
with (OUT / 'D_Develop_Plans/ClassMic_Schedule.csv').open('w', newline='', encoding='utf-8-sig') as f:
    writer = csv.writer(f)
    writer.writerow(['ID', 'Activity', 'Start', 'Finish', 'Calendar days', 'Predecessors', 'Owner', 'On controlling path'])
    writer.writerows(schedule)

prd = '''# ClassMic Product Requirements Document

Status: proposed for team and instructor review. No approval or test result is implied.

## Context and objective

ClassMic tests whether a student phone can make a question easier to hear through a room speaker while the instructor controls speaking turns. The course project covers five user stories in two weekly sprints, followed by evaluation and closing. A live classroom pilot needs a separate decision.

## Target users

- Student speaker: wants to request a turn and be heard without installing a new app.
- Student listener: wants to hear the question clearly.
- Instructor: wants to approve, mute and end speaking without losing control of the class.
- Project manager: needs acceptance results and feedback to support a recommendation.

These are provisional role descriptions, not interview-based personas.

## Requirements and user stories

Import ClassMic_Notion_Requirements.csv as a Notion database. It has the eight fields required by the lab. Keep the requirement ID at the start of Feature/Requirement. Add links to actual evidence and test results after import. Keep Status as Proposed until the appropriate review occurs.

'''
for r in requirements:
    prd += f"### {r['id']} {r['feature']}\n\n{r['story']}\n\nPriority: {r['priority']}. Charter objectives: {r['objective']}. Planned sprint: {r['sprint']}.\n\nAcceptance criteria: {r['acceptance']}\n\nEvidence basis: {r['evidence']}.\n\n"
prd += '''## Priorities and exclusions

R1–R5 are Must. Simple connection checks and a short guide are Should refinements. Visual polish is Could. Recording, saved transcripts, permanent profiles, grading, attendance, captions, polling, simultaneous speakers, remote participation and live campus deployment are excluded.

## Charter objectives

O1: five priority stories and all Must checks, with no open critical defect.
O2: 9/10 setup trials within two minutes and 27/30 ready-speaker activation trials audible within three seconds of approval.
O3: at least 75% positive audibility feedback from at least five willing reviewers, with counts and sample limitations.
O4: no unauthorized transmission in misuse tests, no saved speech and an alternative speaking route.

## Open questions

Obtain formal course acceptance of A2/B2/C1 and D review by 19 September. Validate the audibility problem and instructor workflow. Check devices, any room access and volunteer availability. The selected test target is 95th percentile one-way delay of no more than 150 milliseconds; retain only the minimal records specified in the evaluation protocol. Add the Notion link and genuine human evidence required by C1, D3 and D6.

## Review and maintenance

Nurtore maintains priorities and decisions. Temiko maintains implementation status and technical evidence. The team reviews the PRD at each sprint review. Material scope or schedule changes go to the course instructor. Record approval only when the review actually occurs. Save an export at each milestone.
'''
(OUT / 'D_Develop_Plans/ClassMic_Notion_PRD.md').write_text(prd)

notes = '''# ClassMic evidence and confirmations

This file records evidence status. It is not a substitute for the required original human material.

## Available inputs

- The provided photograph of the in-class SMART note describes ClassMic, one active phone microphone, existing equipment, 75% positive feedback and a proposed 14-week project with a six-week pilot. It is an individual planning artifact, not a photograph of a team charter discussion.
- The earlier local charter and integrated plan supply proposed capabilities and risks. They contain no completed approval record.
- The LMS B page explicitly allows changing the project proposal. Its visible submitted file is the old CampusCrew business case. An upload does not establish approval of the revised ClassMic case.
- The LMS C1 brief permits a draft with clearly marked items requiring confirmation.
- The current D attachment requires two weekly construction sprints, a Notion PRD, at most 20 WBS elements, and at most eight risks including two opportunities.

## Human material still needed or to confirm as accepted separately

1. Genuine photographs or screenshots of the team’s charter brainstorming, with actual participants, date and decisions.
2. The team’s hand-drawn or Post-it WBS and dependency network for D3.
3. Communication brainstorming written before AI prompting for D6.
4. Actual stakeholder comments or observations supporting the requirements. Proposed requirements are labelled until validation occurs.

Individual student inputs are handled separately at the user’s request. No replacement individual contribution has been written or attributed to either student.

## Decisions still needed

- Team and instructor approval of the revised course-prototype scope, A2, B2 and charter.
- Instructor name and acceptance authority, team role assignments, available effort and resources.
- Early D review by 19 September and permission for two weekly sprints on 20–26 September and 27 September–3 October. The E upload area opens on 26 September. A later construction start does not fit two full weeks before the 4 October 00:00 deadline.
- Whether the reviewed C1 is the charter input labelled C2 in D.
- Volunteer reviewers, acceptable audio-delay threshold, test protocol and handling of minimal test metadata.
- Notion page link and access. The PRD and database content are prepared for import but are not hosted in Notion yet.

## Earlier institutional pilot proposal

The earlier documents and SMART note proposed 14 weeks, a six-week pilot, three classrooms and a ceiling of AED 60,000. The old detailed estimate was AED 59,800. These were unverified planning assumptions. This revision uses a two-sprint course prototype as its working scope to match the current D brief. The larger pilot remains a possible later project. It would need new approval, an evidence-based estimate, a commercial comparison and named support and control owners.

## Sources reviewed on 10 September 2026

- B: https://lms.mbzuai.ac.ae/mod/assign/view.php?id=37861
- C: https://lms.mbzuai.ac.ae/mod/assign/view.php?id=38960
- D: https://lms.mbzuai.ac.ae/mod/assign/view.php?id=38962 and D Phase Labs.docx
- E: https://lms.mbzuai.ac.ae/mod/assign/view.php?id=38965
- F: https://lms.mbzuai.ac.ae/mod/assign/view.php?id=38968
- Biamp product overview: https://www.biamp.com/products/families/crowd-mics
- Biamp FAQ: https://www.biamp.com/products/families/crowd-mics/faq
- Biamp design guide: https://support.biamp.com/CrowdMics/Crowd_Mics_design_guide
- W3C Media Capture and Streams: https://www.w3.org/TR/mediacapture-streams/

The assigned textbook passages were not used as evidence of project-specific facts. The documents do not claim completed interviews, research results, testing, approvals or purchased resources.
'''
(OUT / 'Supporting_Notes/Evidence_and_Confirmations.md').write_text(notes)
(ROOT / '.submission_build/content.json').write_text(json.dumps({'objectives': objectives, 'requirements': requirements, 'schedule': schedule, 'effort': effort}, indent=2))
print('Created four DOCX files, Notion PRD, requirements CSV, schedule CSV and evidence record.')
