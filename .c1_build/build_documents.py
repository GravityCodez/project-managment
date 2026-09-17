from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT

ROOT = Path('/Users/gravitycodez/Desktop/Project managment')
OUT = ROOT / 'C1_ClassMic'

def base(title, footer):
    d = Document()
    for border in list(d.styles.element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
    s = d.sections[0]
    s.page_width, s.page_height = Inches(8.5), Inches(11)
    s.top_margin = s.bottom_margin = Inches(.62)
    s.left_margin = s.right_margin = Inches(.7)
    s.header_distance = s.footer_distance = Inches(.27)
    for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2']:
        st = d.styles[name]
        st.font.name = 'Arial'
        st.font.color.rgb = RGBColor(0, 0, 0)
        st.paragraph_format.line_spacing = 1.04
    d.styles['Normal'].font.size = Pt(11)
    d.styles['Normal'].paragraph_format.space_after = Pt(5)
    d.styles['Title'].font.size = Pt(24)
    d.styles['Title'].font.bold = True
    d.styles['Title'].paragraph_format.space_after = Pt(4)
    d.styles['Subtitle'].font.size = Pt(12)
    d.styles['Subtitle'].font.italic = False
    d.styles['Subtitle'].paragraph_format.space_after = Pt(5)
    for name in ['Heading 1', 'Heading 2']:
        d.styles[name].font.size = Pt(11.5)
        d.styles[name].font.bold = True
        d.styles[name].paragraph_format.space_before = Pt(8)
        d.styles[name].paragraph_format.space_after = Pt(4)
        d.styles[name].paragraph_format.keep_with_next = True
    f=s.footer.paragraphs[0]
    f.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    r=f.add_run(footer+'    '); r.font.size=Pt(8.5)
    field=OxmlElement('w:fldSimple'); field.set(qn('w:instr'),'PAGE'); f._p.append(field)
    d.core_properties.title=title
    d.core_properties.subject='C1 Develop the Project Charter'
    d.core_properties.author=''
    d.core_properties.last_modified_by=''
    return d

def p(d, text, bold_prefix=None, size=None):
    x=d.add_paragraph()
    if bold_prefix and text.startswith(bold_prefix):
        x.add_run(bold_prefix).bold=True
        x.add_run(text[len(bold_prefix):])
    else: x.add_run(text)
    if size:
        for r in x.runs: r.font.size=Pt(size)
    return x

def h(d,text): d.add_paragraph(text,'Heading 1')

def table(d,headers,rows,widths,font=10.5):
    t=d.add_table(rows=1,cols=len(headers))
    t.alignment=WD_TABLE_ALIGNMENT.CENTER
    t.autofit=False
    for i,w in enumerate(widths): t.columns[i].width=Inches(w)
    for i,text in enumerate(headers): t.rows[0].cells[i].text=text
    for row in rows:
        for c,text in zip(t.add_row().cells,row): c.text=text
    borders=OxmlElement('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        e=OxmlElement('w:'+edge); e.set(qn('w:val'),'single'); e.set(qn('w:sz'),'4'); e.set(qn('w:color'),'D9D9D9'); borders.append(e)
    t._tbl.tblPr.append(borders)
    for ri,row in enumerate(t.rows):
        trPr=row._tr.get_or_add_trPr()
        no=OxmlElement('w:cantSplit');trPr.append(no)
        if ri==0:
            rep=OxmlElement('w:tblHeader'); trPr.append(rep)
        for ci,cell in enumerate(row.cells):
            cell.width=Inches(widths[ci]); cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            props=cell._tc.get_or_add_tcPr()
            shade=OxmlElement('w:shd'); shade.set(qn('w:fill'),'263F50' if ri==0 else ('F1F4F6' if ri%2 else 'FFFFFF')); props.append(shade)
            margins=OxmlElement('w:tcMar')
            for side,val in [('top',65),('bottom',65),('left',95),('right',95)]:
                e=OxmlElement('w:'+side);e.set(qn('w:w'),str(val));e.set(qn('w:type'),'dxa');margins.append(e)
            props.append(margins)
            for par in cell.paragraphs:
                par.paragraph_format.space_after=Pt(0); par.paragraph_format.line_spacing=1.0
                for r in par.runs:
                    r.font.name='Arial';r.font.size=Pt(font);r.font.bold=ri==0
                    r.font.color.rgb=RGBColor.from_string('FFFFFF' if ri==0 else '000000')
    return t

d=base('ClassMic Project Charter','ClassMic C1 draft')
d.add_paragraph('ClassMic Project Charter','Title')
d.add_paragraph('C1 — Develop the Project Charter','Subtitle')
p(d,'Team  Temiko Machavariani and Nurtore Arynuruly',size=10.5)
p(d,'Draft for confirmation. This charter requests authorization for the classroom pilot. Approved ClassMic A2 and B2 outputs and sponsor approval have not been supplied.',bold_prefix='Draft for confirmation.',size=10.5)
h(d,'1 Purpose and justification')
p(d,"ClassMic aims to make student questions easier to hear and support inclusive classroom participation. Students use their phones as temporary, instructor-controlled microphones, avoiding the need to pass a physical microphone. A limited pilot will test classroom value and technical feasibility using existing Wi-Fi and speakers before any wider investment. The approved aspiration, alternatives appraisal and investment justification require confirmation from ClassMic A2 and B2.")
h(d,'2 Measurable objectives and success criteria')
p(d,'Complete the prototype, testing and six-week pilot in three classrooms by Week 14, within the proposed AED 60,000 ceiling. Evaluate these existing draft targets:')
table(d,['Measure','Target by pilot completion','Evidence'],[
    ['Session setup','At least 90% start within 2 minutes','Timed session records'],
    ['Speaking access','At least 90% of approved requests audible within 3 seconds','Request and audio timing'],
    ['Audio delay','Median delay of 700 ms or less','Latency tests'],
    ['Audibility','At least 75% positive feedback','Student and instructor survey'],
    ['Privacy and security','Zero critical incidents','Incident and control records']
],[1.30,3.75,2.05])
p(d,'Requires confirmation: targets, measurement definitions, minimum sample and acceptance owner.',bold_prefix='Requires confirmation:',size=10)
h(d,'3 High level scope and exclusions')
p(d,'In scope. Browser and QR join, temporary sessions, speaking queue, instructor approval and mute/end controls, one active push-to-talk microphone, classroom speaker output, user guidance, three-room pilot and evaluation.',bold_prefix='In scope.')
p(d,'Excluded. Recording, stored audio or transcripts, voice identification, speech-based attendance or grading, simultaneous microphones, permanent student profiles, remote-learning use and campus-wide rollout.',bold_prefix='Excluded.')
h(d,'4 Deliverables and milestones')
table(d,['Timing','Deliverable or decision'],[
    ['Week 1','Charter approval and kickoff'],
    ['Weeks 2–3','Confirmed classroom, privacy, accessibility and technical requirements'],
    ['Weeks 4–6','Working prototype with instructor controls'],
    ['Weeks 7–8','AV, network, browser, misuse and user acceptance test evidence'],
    ['Weeks 9–14','Six-week classroom pilot, final evaluation and sponsor decision']
],[1.30,5.80],font=10.25)

d.add_page_break()
h(d,'5 High level requirements')
p(d,'Students join a temporary browser session and explicitly permit microphone access. The instructor approves speakers and can mute or end the session. Only one approved phone transmits while push to talk is active. The system sends audio to room speakers without recording or retaining audio or transcripts. Supported devices, accessible participation and a physical microphone or instructor-repeat fallback must be available.')
h(d,'6 Initial risks assumptions and constraints')
p(d,'Risks. Feedback and network delay may make audio unusable. Browser incompatibility or unequal phone access may exclude participants. Misuse or unauthorized audio handling may compromise privacy. Low instructor adoption may leave insufficient evaluation evidence.',bold_prefix='Risks.')
p(d,'Responses. Test rooms and devices before the pilot, use one active microphone and instructor controls, provide training and a fallback, and stop affected sessions for suspected security or privacy incidents.',bold_prefix='Responses.')
p(d,'Assumptions. Stable Wi-Fi, a computer connected to speakers, compatible phones, participating instructors and access to IT, AV, privacy and accessibility reviewers are available. These assumptions require validation.',bold_prefix='Assumptions.')
p(d,'Constraints. The existing draft proposes 14 weeks, three equipped classrooms and an AED 60,000 ceiling. No audio or transcript storage is permitted. Funding, room allocation, dates and resource availability require confirmation.',bold_prefix='Constraints.')
h(d,'7 Sponsor and major stakeholders')
p(d,'Proposed sponsor: Dean of Academic Affairs, with the named sponsor and approval to be confirmed. Major stakeholders are the Teaching and Learning lead as business owner, pilot instructors, students, IT and classroom AV, cybersecurity and privacy, accessibility services, and Finance and Procurement. They confirm classroom value, readiness, controls and resource commitments. Temiko Machavariani is the technical lead named in the existing draft.')
h(d,'8 Project manager and authority')
p(d,'Nurtore Arynuruly is the project manager named in the existing draft, subject to confirmation. On approval, the project manager may coordinate daily work, assign tasks and manage issues within the agreed scope, schedule and ceiling. The sponsor approves material changes and funding exceptions. Relevant IT, AV, privacy and security owners approve readiness in their areas. No authority is effective until the sponsor approves this charter.')
h(d,'Items requiring confirmation')
p(d,'1. Supply the approved ClassMic A2 aspiration and B2 business case. The available CampusCrew business case concerns a different project and cannot establish ClassMic approval.',size=10.5)
p(d,'2. Confirm the sponsor, project manager and technical lead appointments, business owner and control reviewers, funding, calendar dates, rooms, participants and evaluation method.',size=10.5)
p(d,'3. Complete each student’s individual input and attach genuine photos of the team’s charter brainstorming, with dates, participants and decisions. Use the separate input forms and Human Evidence document in this package.',size=10.5)
p(d,'Sponsor name and signature  ________________________    Date  __________',size=10.5)
p(d,'Project manager signature  _________________________    Date  __________',size=10.5)
d.save(OUT/'02_Group_Charter'/'ClassMic_C1_Project_Charter.docx')

labels=[
    'Purpose and justification',
    'Measurable objectives and success criteria',
    'High level scope and exclusions',
    'Deliverables and milestones',
    'High level requirements',
    'Initial risks assumptions and constraints',
    'Sponsor and major stakeholders',
    'Project manager and authority'
]
for name,file in [('Temiko Machavariani','Temiko_Machavariani'),('Nurtore Arynuruly','Nurtore_Arynuruly')]:
    d=base('ClassMic Individual Student Input','ClassMic C1 individual input form')
    d.add_paragraph('ClassMic Individual Student Input','Title')
    p(d,'Student  '+name)
    p(d,'Student ID  __________________    Date  __________________')
    p(d,'To be completed by the student. No individual contribution was supplied. Record your own suggestions and identify the evidence you used from the approved A2 and B2 outputs.',bold_prefix='To be completed by the student.')
    for i,label in enumerate(labels,1):
        h(d,str(i)+' '+label)
        p(d,'[Enter your input and supporting evidence.]',size=10.5)
    h(d,'Team discussion and review')
    p(d,'My proposed changes and the team’s response  [Complete after discussion.]',size=10.5)
    p(d,'Items I still need to confirm  [Enter unresolved questions.]',size=10.5)
    p(d,'Brainstorming photo filename or evidence reference  __________________',size=10.5)
    d.save(OUT/'01_Individual_Inputs'/(file+'_Input_TO_COMPLETE.docx'))

d=base('ClassMic Charter Brainstorming Evidence','ClassMic C1 human evidence form')
d.add_paragraph('ClassMic Charter Brainstorming Evidence','Title')
p(d,'Human evidence required for C1')
p(d,'No brainstorming images were supplied. Insert genuine photographs or screenshots of the team discussing the charter. Add one entry for each discussion image. The fields below are placeholders, not a record of a completed discussion.')
h(d,'Team and discussion details')
p(d,'Team  Temiko Machavariani and Nurtore Arynuruly')
p(d,'Date and location or meeting platform  ______________________________')
p(d,'Participants shown  ____________________________________________')
p(d,'Image filename  _______________________________________________')
h(d,'Brainstorming image')
q=p(d,'[Insert the original discussion photograph or screenshot here.]')
q.paragraph_format.space_after=Pt(132)
h(d,'Evidence caption')
p(d,'Charter sections discussed  [Enter the relevant sections.]')
p(d,'Suggestions from each student  [Describe actual contributions.]')
p(d,'Agreed changes and unresolved items  [Record the team’s decisions.]')
p(d,'Repeat this entry for additional discussion images. Include the originals in this folder when completing the package.',size=10.5)
d.save(OUT/'04_Human_Evidence'/'ClassMic_Brainstorming_Evidence_TO_COMPLETE.docx')

readme='''C1 — Develop the Project Charter
ClassMic draft package

CONTENTS
01_Individual_Inputs: One named form per student. Both require the students’ actual input.
02_Group_Charter: The concise draft charter and items requiring confirmation.
03_Slides: Three slides for a team presentation.
04_Human_Evidence: A form for genuine brainstorming images and discussion captions.

BEFORE UPLOADING
1. Complete the individual input forms using each student’s own contributions.
2. Insert genuine team brainstorming images and complete their captions. Include the original images.
3. Supply the approved ClassMic A2 aspiration and B2 business case. The available Phase B file describes CampusCrew, which is a different project.
4. Confirm the draft charter’s appointments, budget, timing, room availability and success measures. Resolve or retain explicitly labeled confirmation items as instructed by your course.
5. Update the three slides if the agreed charter changes. Re-create the ZIP from this folder after making changes.

The original ClassMic charter supplies the proposed scope, objectives, roles, 14-week schedule, three-classroom pilot and AED 60,000 ceiling. These are carried forward as draft information, without claiming that approval was verified. No student contributions, discussion photos, quotations, approval signatures or missing business-case results have been invented.

LAB READING AS PROVIDED IN THE C1 INSTRUCTIONS
Chapter 5, “AI-Assisted Project Initiation,” pages 147–155.
Chapter 2, selected stakeholder sections, pages 37–53.
The textbook passages were not supplied. These references identify the assigned reading and do not claim that the passages were reviewed.

REQUIRED SUBMISSION
The group charter must remain within two pages. The presentation must contain no more than three slides. Upload one ZIP containing the individual student inputs, group charter and three slides, together with the required human evidence.
'''
(OUT/'START_HERE.txt').write_text(readme)
print('Created charter, two individual input forms, human evidence form and package guide.')
